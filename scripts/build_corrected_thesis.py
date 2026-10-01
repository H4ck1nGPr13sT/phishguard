#!/usr/bin/env python3
"""Build ONE coherent corrected DOCX of the thesis (addressing review 2: K1-K10).

Merges into a single document in standard order — title page first, chapters
1-4, ONE corrected Chapter 5, ONE rewritten conclusion, bibliography, appendices.
The old Chapter 5 / old conclusions and the errata are NOT in the thesis (the
errata is written separately to reports/errata.md). Figures cannot be recovered
from text extraction; the architecture diagram is a cropped render and the UI
screenshots are the real embedded images.

Pipeline: pdftotext -layout -> slice by chapter -> clean + inject fixes ->
Markdown -> pandoc -> python-docx (WSZiB styles + structure).
"""
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC_TXT = Path("/tmp/praca_txt.txt")
OUT_MD = ROOT / "reports/praca-poprawiona.md"
OUT_DOCX = ROOT / "praca-inzynierska-poprawiona.docx"
ERRATA_MD = ROOT / "reports/errata.md"

AUTHOR = "Łukasz Drążek"
PROMOTER = "dr hab. inż. Rafał Dreżewski"

# Chapter boundaries in /tmp/praca_txt.txt (1-indexed, from recon):
TITLE = (215, 266)       # strona tytułowa
CH1_4 = (348, 1860)      # rozdziały 1-4 (body)
BIB = (2405, None)       # 7. Bibliografia -> koniec
# old Chapter 5 (1861-2264) and old conclusions (2265-2404) are dropped.

STALE_NUMBERS = ["0,9701", "0,9747", "0,9764", "0,9778", "0,9686", "0,9733",
                 "0,9751", "0,9766", "0,9840", "0,9823", "0,9685",
                 "+0,0535", "+0,0685", "98,4", "97,47"]

INLINE = [
    (r"\[Imię i Nazwisko Autora\]", AUTHOR),
    (r"\[stopień\.\s*imię nazwisko promotora\]", PROMOTER),
    # U6 + K9 — grammatical whole-sentence fix: three decision layers at
    # prediction time; GA is prior optimization, not a 4th vote. Impersonal.
    (r"Postanowiłem zaprojektować i wykonać system wykrywania phishingu, w którym wynik\s+"
     r"klasyfikacji opiera się nie na jednym modelu, ale na połączonej analizie czterech uzupełniających\s+"
     r"się podejść: uczenia maszynowego, algorytmu ewolucyjnego, systemu regułowego i",
     "Zaprojektowano i wykonano system wykrywania phishingu, w którym wynik "
     "klasyfikacji opiera się nie na jednym modelu, lecz na trzech warstwach "
     "decyzyjnych działających podczas predykcji: zespole uczenia maszynowego, "
     "systemie regułowym i klasyfikatorze bayesowskim. Algorytm genetyczny pełni "
     "rolę wcześniejszej optymalizacji modeli, a nie czwartej warstwy głosującej. "
     "Trzy warstwy predykcji łączą: uczenie maszynowe, system regułowy i podejście"),
    # U7 — Optuna default sampler is TPE, not a Gaussian process.
    (r"Ma jednak wady\.\s*Modeluje funkcję jako proces gaussowski,\s*co przy\s*"
     r"zmiennych kategorialnych wymaga sztucznych przekształceń\.\s*Poza tym może utykać w\s*"
     r"lokalnych ekstremach\.",
     "Domyślnym samplerem biblioteki Optuna jest TPE (Tree-structured Parzen "
     "Estimator), a wariant oparty na procesie Gaussa stanowi osobną, opcjonalną "
     "metodę. Optymalizacja bayesowska bywa jednak wrażliwa na zmienne kategorialne "
     "i dobór przestrzeni; w tym projekcie wartości kategorialne i tak mapowane są "
     "na indeksy (src/optimization/search_spaces.py). Algorytm genetyczny wybrano ze "
     "względu na naturalną obsługę mieszanych przestrzeni oraz prostą kontrolę "
     "kosztu obliczeń, a nie z powodu udowodnionej przewagi nad optymalizacją "
     "bayesowską."),
    # K9 — impersonal form.
    (r"Do projektu wybrałem", "Do projektu wybrano"),
    (r"Optymalizację zrealizowałem", "Optymalizację zrealizowano"),
    (r"Tabela 0\.", "Tabela 1."),
]

RESULTS = """
# 5. Wyniki projektu

Wszystkie wyniki URL pochodzą z jednego protokołu: wspólny odłożony zbiór testowy
z artefaktu `cache/url_training_data.joblib` (200 próbek treningowych, 50 testowych,
25/25, 30 cech), podział **losowy stratyfikowany** (`random_state=42`) — nie
temporalny. Skrypt: `scripts/evaluate_url_models_honest.py`, raport:
`reports/url_eval_honest.csv`. Mały zbiór testowy (50 próbek) oznacza, że różnice
są orientacyjne i opisowe.

## 5.1. Klasyfikatory bazowe i zoptymalizowane GA (wspólny test)

Każdy klasyfikator oceniono na tym samym odłożonym teście. Kolumna „CV-fitness" to
F1 z 5-krotnej walidacji krzyżowej najlepszego osobnika GA i jest inną wielkością
niż wynik testowy; nie należy ich od siebie odejmować.

| Klasyfikator | F1 baza (test) | F1 GA (test) | Zysk (test) | CV-fitness |
|---|---|---|---|---|
| Logistic Regression | 0,9200 | 0,9362 | +0,0162 | 0,9749 |
| Random Forest | 0,9091 | 0,9259 | +0,0168 | 0,9705 |
| MLP | 0,9167 | 0,9583 | +0,0417 | 0,9697 |
| XGBoost | 0,9231 | 0,9259 | +0,0028 | 0,9623 |
| SVM (RBF) | 0,8846 | 0,9020 | +0,0173 | 0,9502 |
| Naive Bayes | 0,9020 | 0,9020 | +0,0000 | 0,9438 |
| Decision Tree | 0,8980 | 0,9259 | +0,0280 | 0,9656 |

Na wspólnym teście optymalizacja GA dała niewielką poprawę F1 (od +0,000 dla Naive
Bayes do +0,042 dla MLP). Jest to wynik jednego losowego podziału; stabilności
poprawy przy innych ziarnach nie badano, a koszt wyszukiwania (populacja 50,
30 generacji) nie jest tu zestawiony z wielkością zysku.

## 5.2. Zespoły klasyfikatorów (wspólny test)

| Wariant | F1 (test) |
|---|---|
| Hard voting | 0,9200 |
| Soft voting | 0,9231 |
| Stacking | 0,9231 |
| Najlepszy pojedynczy (MLP-GA) | 0,9583 |

Na wspólnym zbiorze testowym zespoły osiągnęły F1 ≈ 0,92 i **nie przewyższyły**
najlepszego pojedynczego modelu zoptymalizowanego GA (MLP, 0,958). W tym
eksperymencie zespół nie uzyskał przewagi; wynik ten opisano zgodnie z pomiarem.

## 5.3. Modele wyspecjalizowane (e-mail, SMS) — dane syntetyczne

Trzy różne wielkości należy rozróżnić. Zapisane w modelach pole `test_accuracy=1,0`
pochodzi z **syntetycznego odłożonego zbioru 20%** utworzonego przez skrypt
treningowy (`scripts/train_email_sms_models.py`, podział 80/20). Niezależne
5-krotne CV (`scripts/evaluate_email_sms_honest.py`) dało: e-mail acc 0,985 /
F1 0,984, SMS acc 1,0 / F1 1,0. Ewaluacja na pełnym zbiorze daje 1,0 (ten sam
zbiór, na którym trenowano) i jest wyłącznie kontrolą, nie miarą generalizacji.

| Typ | 5-fold CV acc | 5-fold CV F1 |
|---|---|---|
| E-mail (65 cech) | 0,985 | 0,984 |
| SMS (70 cech) | 1,000 | 1,000 |

Istotne ograniczenie zakresu: dane są syntetyczne i generowane z 20 szablonów na
klasę dla każdego typu (`scripts/create_email_sms_dataset.py`). W walidacji
krzyżowej większość wiadomości testowych ma swój szablon obecny w zbiorze
treningowym, dlatego CV mierzy przede wszystkim rozpoznawanie **wariantów znanych
szablonów**, a nie nowych kampanii ani rzeczywistych wiadomości. Wynik SMS równy
1,0 wynika z trywialnej separowalności tych danych. Twierdzenie o skuteczności
poza danymi syntetycznymi wymagałoby podziału grupowego według szablonu albo
niezależnego, ręcznie oznaczonego zbioru rzeczywistych wiadomości.

## 5.4. System wielowarstwowy — miara rozbieżności, wkład warstw, czas

**Miara rozbieżności.** Wskaźnik rozbieżności to znormalizowana entropia Shannona
rozkładu głosów, dzielona przez maksymalną entropię binarną log2(2) = 1 bit
(`src/models/disagreement.py`, `src/paradigms/aggregation/disagreement.py`;
17 testów jednostkowych). Dla trzech warstw każda niejednomyślność daje wynik
≈ 0,918, a jednomyślność 0. Dla siedmiu klasyfikatorów najbardziej wyrównany
podział 4/3 daje maksimum ≈ 0,985. Próg oznaczania przypadku granicznego ustawiono
na 0,7: dla warstw oznacza to „paradygmaty nie są jednomyślne", a dla zespołu ML —
„co najmniej dwa z siedmiu klasyfikatorów są odmiennego zdania". Flaga jest
wskaźnikiem niezgody głosów, a nie dowodem, że dana decyzja jest błędna.

**Wkład warstw (próba ilustracyjna z nakładaniem danych).** Na zbiorze 50
etykietowanych adresów (40 phishingowych z OpenPhish, 10 znanych legalnych) sam
zespół ML osiągnął 48/50 poprawnych decyzji, a pełny agregator trójwarstwowy 46/50.
Agregator zmienił dwie decyzje względem samego ML i obie okazały się błędne. Jest
to jednak **próba ilustracyjna, nie niezależna ocena skuteczności**: pierwsze 40
adresów pochodzi z tego samego pliku OpenPhish, z którego losowano dane treningowe
URL — po odtworzeniu losowania 13 z 40 trafia do zbioru treningowego, a 4 do
testowego. Dla rzetelnej oceny wkładu warstw potrzebna byłaby niezależna próba
(z deduplikacją) z pełnym raportem zmian decyzji oraz błędów FN/FP. W badanym
zakresie integracja trzech warstw nie poprawiła decyzji względem samego ML.

**Czas odpowiedzi.** Pomiar `/predict/multi-paradigm` (50 żądań, model rozgrzany,
cechy leksykalne, pomiar lokalny in-process przez TestClient): mediana ≈ 15 ms,
95. percentyl ≈ 16 ms — poniżej wymaganych 500 ms. Zakres twierdzenia ograniczono
do pojedynczego żądania po rozgrzaniu; nie obejmuje ono opóźnienia klient–serwer
ani obciążenia równoległego.

"""

CONCLUSION = """
# 6. Podsumowanie

Zaprojektowano i wykonano system PhishGuard łączący ekstrakcję cech, siedem
klasyfikatorów uczenia maszynowego, optymalizację hiperparametrów algorytmem
genetycznym, system regułowy, klasyfikator bayesowski oraz warstwę agregacji z
interfejsem API i demonstratorem. Wkład własny obejmuje integrację komponentów,
dobór cech i mechanizm wykrywania rozbieżności między warstwami.

Wyniki należy odczytywać w zakresie przeprowadzonych pomiarów. Na wspólnym,
losowo podzielonym zbiorze testowym 50 adresów URL optymalizacja GA dała niewielką
poprawę F1 (do +0,042 dla MLP), a zespoły klasyfikatorów nie przewyższyły
najlepszego pojedynczego modelu zoptymalizowanego GA. Wyniki e-mail i SMS uzyskano
na danych syntetycznych generowanych z ograniczonej liczby szablonów i nie mierzą
one skuteczności na rzeczywistych wiadomościach. Na ilustracyjnej próbie 50 adresów
agregator trójwarstwowy uzyskał 46/50 poprawnych decyzji wobec 48/50 dla samego
zespołu ML; nie wykazano wzrostu skuteczności wynikającego z integracji warstw.
Naprawiono natomiast główny mechanizm pracy: po poprawieniu normalizacji entropii
flaga rozbieżności jest osiągalna i poprawnie sygnalizuje niezgodę głosów.
Czas odpowiedzi pojedynczego żądania po rozgrzaniu modelu mieści się poniżej 500 ms
w pomiarze lokalnym.

System pokazuje wyniki trzech warstw oraz aktywne reguły, co zwiększa
interpretowalność decyzji. Jest to cecha interfejsu, odrębna od zmierzonej
skuteczności klasyfikacji. Dalsze prace, które pozwoliłyby rozszerzyć wnioski, to
w szczególności: ewaluacja na niezależnych, rzeczywistych zbiorach wiadomości i
adresów, podział grupowy danych syntetycznych według szablonu, zestawienie kosztu
wyszukiwania GA z uzyskanym zyskiem oraz powtórzenia optymalizacji dla oceny jej
stabilności. Zakres opcjonalny obejmuje porównanie z modelem głębokim oraz
rozszerzenia OCR i wyjaśnień SHAP/LIME. Praca nie obejmowała bezpośredniego
porównania z modelami głębokimi i nie formułuje wniosku o przewadze nad nimi.

"""

FIGURES = [
    ("reports/figury2/diagram_architektura.png", "Rysunek A1. Architektura logiczna systemu PhishGuard (przycięty diagram, s. 25)."),
    ("reports/figury2/zrzut-000.png", "Rysunek A2. Interfejs — zrzut ekranu (s. 48 oryginału)."),
    ("reports/figury2/zrzut-001.png", "Rysunek A3. Interfejs — zrzut ekranu (s. 49 oryginału)."),
    ("reports/figury2/zrzut-002.png", "Rysunek A4. Interfejs — zrzut ekranu (s. 50 oryginału)."),
]


def clean(lines):
    out = []
    for ln in lines:
        s = ln.rstrip()
        if re.fullmatch(r"\s*\d{1,3}\s*", s):
            continue
        s = re.sub(r"\.{4,}\s*", " — ", s)
        if " — " not in s and any(tok in s for tok in STALE_NUMBERS):
            continue
        if any(ch in s for ch in "│┌┐└┘├┤┬┴┼"):
            continue
        toks = s.split()
        if len(toks) >= 6 and sum(1 for t in toks if len(t) == 1) / len(toks) > 0.5:
            continue
        out.append(s)
    return "\n".join(out)


def _slice(all_lines, a, b):
    return all_lines[a - 1: (b if b else len(all_lines))]


def main():
    if not SRC_TXT.exists():
        print("Brak /tmp/praca_txt.txt — uruchom pdftotext -layout najpierw.")
        return 1
    L = SRC_TXT.read_text(encoding="utf-8", errors="replace").splitlines()

    title = clean(_slice(L, *TITLE))
    body = clean(_slice(L, *CH1_4))
    bib = clean(_slice(L, *BIB))
    for pat, rep in INLINE:
        title = re.sub(pat, rep, title)
        body, _ = re.subn(pat, rep, body)

    md = (title + "\n\n" + body + "\n\n" + RESULTS + "\n" + CONCLUSION + "\n"
          + "# 7. Bibliografia\n\n" + bib + "\n")
    # Appendix A — figures
    md += "\n\n# Załącznik A — rysunki\n\n"
    for rel, cap in FIGURES:
        p = ROOT / rel
        if p.exists():
            md += f"![{cap}]({p})\n\n*{cap}*\n\n"
    # Appendix B — clean code listings
    md += "\n\n# Załącznik B — wybrane listingi kodu\n\n"
    for path, t in [
        ("src/paradigms/aggregation/disagreement.py", "Miara rozbieżności (poprawiona normalizacja)"),
        ("src/paradigms/aggregation/weights.py", "Wagi paradygmatów (ML 0,5 / reguły 0,3 / Bayes 0,2)"),
    ]:
        p = ROOT / path
        if p.exists():
            md += f"**{t}** (`{path}`):\n\n```python\n{p.read_text(encoding='utf-8')[:2200]}\n```\n\n"
    OUT_MD.write_text(md, encoding="utf-8")

    ERRATA_MD.write_text(
        "# Errata — wykaz poprawek (dla promotora, poza egzemplarzem pracy)\n\n"
        "Scalono do jednego dokumentu: strona tytułowa, rozdz. 1-4, jeden rozdz. 5, "
        "jedno zakończenie, bibliografia, załączniki. Usunięto stary rozdz. 5 i stare "
        "wnioski. Poprawki merytoryczne: U1 (rozbieżność), U2/U3 (wspólny test URL), "
        "U4 (e-mail/SMS: odłożony 20% syntetyczny + CV, przeciek szablonów), U9/U10 "
        "(czas, wkład warstw z nakładaniem danych oznaczony jako ilustracyjny), "
        "U6/U7 (trzy warstwy + GA, Optuna=TPE), K9 (forma bezosobowa).\n", encoding="utf-8")

    subprocess.run(["pandoc", str(OUT_MD), "-o", str(OUT_DOCX)], check=True)
    _apply_wszib_styles(OUT_DOCX)
    _apply_structure(OUT_DOCX)
    print(f"OK -> {OUT_DOCX.relative_to(ROOT)} ({OUT_DOCX.stat().st_size} B)")
    return 0


def _apply_wszib_styles(path):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = Document(str(path))
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)
    n = d.styles["Normal"]
    n.font.name = "Times New Roman"; n.font.size = Pt(12)
    n.paragraph_format.line_spacing = 1.5
    n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12), ("Title", 18)]:
        try:
            st = d.styles[name]
            st.font.name = "Times New Roman"; st.font.size = Pt(size); st.font.bold = True
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        except KeyError:
            pass
    d.save(str(path))


def _field(run, instr, placeholder):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = placeholder
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (b, it, sep, t, end):
        run._r.append(el)


def _apply_structure(path):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = Document(str(path))
    h1 = re.compile(r"^\d+\.\s+\S")
    h2 = re.compile(r"^\d+\.\d+\.?\s+\S")
    h3 = re.compile(r"^\d+\.\d+\.\d+\.?\s+\S")
    cap = re.compile(r"^(Rysunek|Tabela|Wykres|Schemat)\s+\d", re.IGNORECASE)
    for p in d.paragraphs:
        txt = p.text.strip()
        if not txt or len(txt) > 90:
            continue
        try:
            if h3.match(txt):
                p.style = d.styles["Heading 3"]
            elif h2.match(txt):
                p.style = d.styles["Heading 2"]
            elif h1.match(txt):
                p.style = d.styles["Heading 1"]
            elif cap.match(txt):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(10); r.font.bold = True
        except KeyError:
            pass
    # Insert the TOC AFTER the title page — before the first "1. Wstęp" heading,
    # so the order is: title page, Spis treści, chapters.
    anchor = next((p for p in d.paragraphs
                   if re.match(r"^1\.\s+Wstęp", p.text.strip())), d.paragraphs[0])
    toc_p = anchor.insert_paragraph_before("Spis treści")
    toc_p.style = d.styles["Heading 1"]
    toc_field_p = toc_p.insert_paragraph_before("")
    toc_p._p.addnext(toc_field_p._p)
    _field(toc_field_p.add_run(), 'TOC \\o "1-3" \\h \\z \\u',
           "Spis treści — kliknij i naciśnij F9, aby zaktualizować")
    for sec in d.sections:
        sec.different_first_page_header_footer = True
        f = sec.footer
        fp = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(fp.add_run(), "PAGE", "1")
        sec.first_page_footer.is_linked_to_previous = False
    d.save(str(path))


if __name__ == "__main__":
    sys.exit(main())
